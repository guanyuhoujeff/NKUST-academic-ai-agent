#!/usr/bin/env python3
"""deck.json → slides.pptx（可編輯的原生文字與表格，中文字型寫進東亞字型槽）
用法: python build.py [deck.json] [slides.pptx] [--pdf]
建完會自動檢查並印出結果：
  [字型] 每一段文字都有東亞字型（<a:ea>）
  [大綱] 每一頁的標題都出自確認過的 outline.md
  [份量] 標題不籠統、每頁要點數與字數沒超過上限、版面有變化
  [語言] 沒有簡體字
  [來源] 投影片裡的數字、DOI、ISSN 都能在 sources 列的檔案裡找到
[大綱]、[語言] 或 [來源] 沒過時結束代碼為 2。加 --pdf 會另存一份 slides.pdf 方便預覽（需要 LibreOffice）。
"""
import json, re, sys, platform, pathlib, shutil, subprocess
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.enum.shapes import MSO_SHAPE

# Windows 繁中版的主控台是 cp950：印到不在 cp950 的字（論文標題常見的「‐」「é」「∑」）會當掉，
# 而且 opencode 會把 cp950 的中文讀成亂碼，agent 看不懂檢查結果。一律改用 UTF-8 輸出。
for _stream in (sys.stdout, sys.stderr):
    try: _stream.reconfigure(encoding="utf-8", errors="replace")
    except Exception: pass

# 中文要正確顯示，必須把字型寫進「東亞字型槽」<a:ea>；只設一般字型名稱的話，中文會掉到新細明體
EA = {"Windows": "Microsoft JhengHei", "Darwin": "PingFang TC"}.get(platform.system(), "Noto Sans CJK TC")
LATIN = "Arial"
MAX_BULLETS, MAX_CHARS, MAX_SAME_LAYOUT = 5, 45, 3
MAX_SECTIONS = 2   # 章節頁最多連續 2 頁（例如「小結」接「Q&A」）
LAYOUTS = ("section", "assertion", "bignumber", "table", "compare", "quote")
# 只由這些詞組成的標題等於沒說話（例：「研究背景與目的」「結論與未來方向」）
GENERIC = re.compile(r"^(研究|本文|文獻|主要|相關|整體|簡報)*(歡迎|議程|簡報|背景|目的|目標|動機|範圍|方法|設計|資料|結果|發現|結論|建議|限制|摘要|概覽|概述|總覽|簡介|介紹|回顧|討論|趨勢|未來方向|展望|與|及|和|、|：|\s)+$")

# 只收「繁體裡不存在」的簡體字（OpenCC 的 STCharacters 扣掉 TSCharacters 與 TWVariants 兩邊出現過的字，共 3796 字）
SIMPLIFIED = set("㐷㐽㑇㑈㑔㑩㓆㓥㓰㔉㖊㖞㘎㚯㛀㛟㛠㛣㛤㛿㟆㟜㟥㡎㤘㤽㥪㧏㧐㧑㧟㧰㨫㭎㭏㭣㭤㭴㱩㱮㲿㳔㳕㳠㳡㳢㳽㴋㶉㶶㶽㺍㻅㻏㻘䀥䁖䂵䃅䅉䅟䅪䇲䉤䌶䌷䌸䌹䌺䌻䌼䌽䌾䌿䍀䍁䍠䎬䏝䑽䓓䓕䓖䓨䗖䘛䘞䙊䙌䙓䜣䜤䜥䜧䜩䝙䞌䞍䞎䞐䟢䢀䢁䢂䥺䥽䥾䥿䦀䦁䦂䦃䦅䦆䦶䦷䩄䭪䯃䯄䯅䲝䲞䲟䲠䲡䲢䲣䴓䴔䴕䴖䴗䴘䴙䶮与专业丛东丝丢两严丧个临为丽举义乌乐乔习乡书买乱争亏亚产亩亲亵亸亿仅从仑仓仪们众优会伛伞伟传伡伣伤伥伦伧伪伫体佥侠侣侥侦侧侨侩侪侬侭俣俦俨俩俪俫俭债倾偬偻偾偿傤傥傧储傩儿兑兖兰关兴兹养兽冁内冈册写军农冯冲决况冻净凄凉减凑凛凤凫凭凯击凿刍刘则刚创删别刬刭刹刽刾刿剀剂剐剑剥剧劝办务劢动励劲劳势勋勚匀匦匮区医华协单卖卢卤卧卫却卺厅历厉压厌厍厐厕厢厣厦厨厩厮县叁参叆叇双发变叙叠号叹叽吓吕吗吨听启吴呐呒呓呕呖呗员呙呛呜咏咙咛咝咤响哑哒哓哔哕哗哙哜哝哟唛唝唠唡唢唤啧啬啭啮啯啰啴啸喷喽喾嗫嗳嘘嘤嘱噜嚣团园囱围囵国图圆圣圹场坏块坚坛坜坝坞坟坠垄垅垆垒垦垩垫垭垯垱垲垴埘埙埚堑堕塆墙壮声壳壶壸处备复够头夹夺奁奂奋奖奥妆妇妈妩妪妫姗姹娄娅娆娇娈娱娲娴婳婴婵婶媪媭嫒嫔嫱嬷孙学孪宁宝实宠审宪宫宽宾寝对寻导寿将尔尘尝尧尴尽层屃屉届属屡屦屿岁岂岖岗岘岚岛岭岽岿峃峄峡峣峤峥峦崂崃崄崭嵘嵚嵝巅巩巯币帅师帏帐帜带帧帮帱帻帼幂并庄庆庐庑库应庙庞废庼廪开异弃弑张弥弪弯弹强归当录彟彦彨彻径徕忆忏忧忾怀态怂怃怄怅怆怜总怼怿恋恒恳恶恸恹恺恻恼恽悦悫悬悭悮悯惊惧惨惩惫惬惭惮惯愠愤愦慑慭懑懒懔戆戋戏戗战戬戯户扑执扩扪扫扬扰抚抛抟抠抡抢护报担拟拢拣拥拦拧拨择挚挛挜挝挞挟挠挡挢挣挤挥挦捝捞损捡换捣掳掴掷掸掺掼揽揾揿搀搁搂搄搅携摄摅摆摇摈摊撄撑撵撷撸撺擜擞攒敌敚敛敩数斋斓斩断无旧时旷旸昙昵昼昽显晋晒晓晔晕晖暂暅暧术机杀杂权条来杨杩构枞枢枣枥枧枨枪枫枭柠柽栀栅标栈栉栊栋栌栎栏树栖样栾桠桡桢档桤桥桦桧桨桩桪梦梼梾梿检棁棂椁椝椟椠椢椤椫椭椮楼榄榅榇榈榉榝槚槛槟槠横樯樱橥橱橹橼檩欢欤欧歼殁殇残殒殓殚殡殴毁毂毕毙毡毵毶氇气氢氩氲汇汉汤汹沄沟没沣沤沥沦沧沨沩沪泞泪泶泷泸泺泻泼泽泾洁洒洼浃浅浆浇浈浉浊测浍济浏浐浑浒浓浔浕涚涛涝涞涟涠涡涢涣涤润涧涨涩渊渌渍渎渐渑渔渖渗温湾湿溁溃溅溆溇滗滚滞滟滠满滢滤滥滦滨滩滪潆潇潋潍潜潴澛澜濑濒灏灭灯灵灾灿炀炉炖炜炝点炼炽烁烂烃烛烟烦烧烨烩烫烬热焕焖焘煴爱爷牍牦牵牺犊状犷犸犹狈狝狞独狭狮狯狰狱狲猃猎猕猡猪猫猬献獭玑玙玚玛玮环现玱玺珐珑珰珲琎琏琐琼瑶瑷瑸璎瓒瓮瓯电画畅畴疖疗疟疠疡疬疭疮疯疱疴痈痉痒痖痨痪痫瘅瘆瘗瘘瘪瘫瘾瘿癞癣癫皑皱皲盏盐监盖盗盘眍眦眬睁睐睑瞆瞒瞩矫矶矾矿砀码砖砗砚砜砺砻砾础硁硕硖硗硙硚硵硷碍碛碜碱礼祃祎祢祯祷祸禀禄禅离秃秆积称秽秾稆税稣稳穑穞穷窃窍窎窑窜窝窥窦窭竖竞笃笋笔笕笺笼笾筚筛筜筝筹筼签筿简箓箦箧箨箩箪箫篑篓篮篯篱簖籁籴类籼粜粝粤粪粮糁糇糍紧絷縆纟纠纡红纣纤纥约级纨纩纪纫纬纭纮纯纰纱纲纳纴纵纶纷纸纹纺纻纼纽纾线绀绁绂练组绅细织终绉绊绋绌绍绎经绐绑绒结绔绕绖绗绘给绚绛络绝绞统绠绡绢绣绤绥绦继绨绩绪绫绬续绮绯绰绱绲绳维绵绶绷绸绹绺绻综绽绾绿缀缁缂缃缄缅缆缇缈缉缊缋缌缍缎缏缐缑缒缓缔缕编缗缘缙缚缛缜缝缞缟缠缡缢缣缤缥缦缧缨缩缪缫缬缭缮缯缰缱缲缳缴缵罂网罗罚罢罴羁羟羡翘翙翚耢耧耸耻聂聋职聍联聩聪肃肠肤肮肴肾肿胀胁胆胧胨胪胫胶脉脍脏脐脑脓脔脚脱脶脸腘腭腻腼腽腾膑臜舆舣舰舱舻艰艳艺节芈芗芜芦苁苇苈苋苌苍苎苏茎茏茑茔茕茧荆荙荚荛荜荝荞荟荠荡荣荤荥荦荧荨荩荪荫荬荭荮药莅莱莲莳莴莶获莸莹莺莼萚萝萤营萦萧萨葱蒀蒇蒉蒋蒌蒏蓝蓟蓠蓣蓥蓦蔂蔷蔹蔺蔼蕰蕲蕴薮藓蘖虏虑虚虬虮虱虽虾虿蚀蚁蚂蚃蚕蚬蛊蛎蛏蛮蛰蛱蛲蛳蛴蜕蜗蝇蝈蝉蝼蝾螀螨蟏衅衔补衬衮袄袅袆袜袭袯装裆裈裢裣裤裥褛褴襕见观觃规觅视觇览觉觊觋觌觍觎觏觐觑觞触觯訚詟誉誊讠计订讣认讥讦讧讨让讪讫讬训议讯记讱讲讳讴讵讶讷许讹论讻讼讽设访诀证诂诃评诅识诇诈诉诊诋诌词诎诏诐译诒诓诔试诖诗诘诙诚诛诜话诞诟诠诡询诣诤该详诧诨诩诪诫诬语诮误诰诱诲诳说诵诶请诸诹诺读诼诽课诿谀谁谂调谄谅谆谇谈谉谊谋谌谍谎谏谐谑谒谓谔谕谖谗谘谙谚谛谜谝谞谟谠谡谢谣谤谥谦谧谨谩谪谫谬谭谮谯谰谱谲谳谴谵谶豮贝贞负贠贡财责贤败账货质贩贪贫贬购贮贯贰贱贲贳贴贵贶贷贸费贺贻贼贽贾贿赀赁赂赃资赅赆赇赈赉赊赋赌赍赎赏赐赑赒赓赔赕赖赗赘赙赚赛赜赝赞赟赠赡赢赣赪赵赶趋趱趸跃跄跞践跶跷跸跹跻踌踪踬踯蹑蹒蹰蹿躏躜躯车轧轨轩轪轫转轭轮软轰轱轲轳轴轵轶轷轸轹轺轻轼载轾轿辀辁辂较辄辅辆辇辈辉辊辋辌辍辎辏辐辑辒输辔辕辖辗辘辙辚辞辩辫边辽达迁过迈运还这进远违连迟迩迳迹选逊递逦逻遗遥邓邝邬邮邹邺邻郏郐郑郓郦郧郸酂酝酦酱酽酾酿释鉴銮錾钅钆钇针钉钊钋钌钍钎钏钐钑钒钓钔钕钖钗钘钙钚钛钜钝钞钟钠钡钢钣钤钥钦钧钨钩钪钫钬钭钮钯钰钱钲钳钴钵钶钷钸钹钺钻钼钽钾钿铀铁铂铃铄铅铆铇铈铉铊铋铌铍铎铏铐铑铒铓铔铕铖铗铘铙铚铛铜铝铞铟铠铡铢铣铤铥铦铧铨铩铪铫铬铭铮铯铰铱铲铳铴铵银铷铸铹铺铻铼铽链铿销锁锂锃锄锅锆锇锈锉锊锋锌锍锎锏锐锑锒锓锔锕锖锗锘错锚锛锜锝锞锟锠锡锢锣锤锥锦锧锨锩锪锫锬锭键锯锰锱锲锳锴锵锶锷锸锹锺锻锼锽锾锿镀镁镂镃镄镅镆镇镈镉镊镋镌镍镎镏镐镑镒镓镔镕镖镗镘镙镚镛镜镝镞镟镠镡镢镣镤镥镦镧镨镩镪镫镬镭镮镯镰镱镲镳镴镵镶长门闩闪闫闬闭问闯闰闱闲闳间闵闶闷闸闹闺闻闼闽闾闿阀阁阂阃阄阅阆阇阈阉阊阋阌阍阎阏阐阑阒阓阔阕阖阗阘阙阚阛队阳阴阵阶际陆陇陈陉陕陦陧陨险随隐隶隽难雇雏雠雳雾霁霉霡霭靓靔静靥鞑鞒鞯鞲韦韧韨韩韪韫韬韵页顶顷顸项顺须顼顽顾顿颀颁颂颃预颅领颇颈颉颊颋颌颍颎颏颐频颒颓颔颕颖颗题颙颚颛颜额颞颟颠颡颢颣颤颥颦颧风飏飐飑飒飓飔飕飖飗飘飙飚飞飨餍饣饤饥饦饧饨饩饪饫饬饭饮饯饰饱饲饳饴饵饶饷饸饹饺饻饼饽饾饿馀馁馂馃馄馅馆馇馈馉馊馋馌馍馎馏馐馑馒馓馔馕马驭驮驯驰驱驲驳驴驵驶驷驸驹驺驻驼驽驾驿骀骁骂骃骄骅骆骇骈骉骊骋验骍骎骏骐骑骒骓骔骕骖骗骘骙骚骛骜骝骞骟骠骡骢骣骤骥骦骧髅髋髌鬓鬶魇魉鱼鱽鱾鱿鲀鲁鲂鲃鲄鲅鲆鲇鲈鲉鲊鲋鲌鲍鲎鲏鲐鲑鲒鲓鲔鲕鲖鲗鲘鲙鲚鲛鲜鲝鲞鲟鲠鲡鲢鲣鲤鲥鲦鲧鲨鲩鲪鲫鲬鲭鲮鲯鲰鲱鲲鲳鲴鲵鲶鲷鲸鲹鲺鲻鲼鲽鲾鲿鳀鳁鳂鳃鳄鳅鳆鳇鳈鳉鳊鳋鳌鳍鳎鳏鳐鳑鳒鳓鳔鳕鳖鳗鳘鳙鳚鳛鳜鳝鳞鳟鳠鳡鳢鳣鳤鸟鸠鸡鸢鸣鸤鸥鸦鸧鸨鸩鸪鸫鸬鸭鸮鸯鸰鸱鸲鸳鸴鸵鸶鸷鸸鸹鸺鸻鸼鸽鸾鸿鹀鹁鹂鹃鹄鹅鹆鹇鹈鹉鹊鹋鹌鹍鹎鹏鹐鹑鹒鹓鹔鹕鹖鹗鹘鹙鹚鹛鹜鹝鹞鹟鹠鹡鹢鹣鹤鹥鹦鹧鹨鹩鹪鹫鹬鹭鹮鹯鹰鹱鹲鹳鹴鹾麦麸麹麺黄黉黡黩黪黾鼋鼌鼍鼹齐齑齿龀龁龂龃龄龅龆龇龈龉龊龋龌龙龚龛龟鿎鿏鿒鿔𠀾𠆲𠆿𠇹𠉂𠉗𠋆𠚳𠛅𠛆𠛾𠡠𠮶𠯟𠯠𠰱𠰷𠱞𠲥𠴛𠴢𠵸𠵾𡋀𡋗𡋤𡍣𡒄𡝠𡞋𡞱𡠟𡥧𡭜𡭬𡳃𡳒𡶴𡸃𡺃𡺄𢋈𢗓𢘙𢘝𢘞𢙏𢙐𢙑𢙒𢙓𢛯𢠁𢢐𢧐𢫊𢫞𢫬𢬍𢬦𢭏𢽾𣃁𣆐𣈣𣍨𣍯𣍰𣎑𣏢𣐕𣐤𣑶𣒌𣓿𣔌𣗊𣗋𣗙𣘐𣘓𣘴𣘷𣚚𣞎𣨼𣭤𣯣𣱝𣲗𣲘𣳆𣶩𣶫𣶭𣷷𣸣𣺼𣺽𣽷𤆡𤆢𤇃𤇄𤇭𤇹𤈶𤈷𤊀𤊰𤋏𤎺𤎻𤙯𤝢𤞃𤞤𤠋𤦀𤩽𤳄𤶊𤶧𤻊𤽯𤾀𤿲𥁢𥅘𥅴𥅿𥆧𥇢𥎝𥐟𥐯𥐰𥐻𥞦𥧂𥩟𥩺𥫣𥬀𥬞𥬠𥭉𥮋𥮜𥮾𥱔𥹥𥺅𥺇𦈈𦈉𦈋𦈌𦈎𦈏𦈐𦈑𦈒𦈓𦈔𦈕𦈖𦈗𦈘𦈙𦈚𦈛𦈜𦈝𦈞𦈟𦈠𦈡𦍠𦛨𦝼𦟗𦨩𦰏𦰴𦶟𦶻𦻕𧉐𧉞𧌥𧏖𧏗𧑏𧒭𧜭𧝝𧝧𧮪𧳕𧹑𧹒𧹓𧹔𧹕𧹖𧹗𧿈𨀁𨀱𨁴𨂺𨄄𨅛𨅫𨅬𨉗𨐅𨐆𨐇𨐈𨐉𨐊𨑹𨟳𨠨𨡙𨡺𨤰𨰾𨰿𨱀𨱁𨱂𨱃𨱄𨱅𨱆𨱇𨱈𨱉𨱊𨱋𨱌𨱍𨱎𨱏𨱐𨱑𨱒𨱓𨱔𨱕𨱖𨷿𨸀𨸁𨸂𨸃𨸄𨸅𨸆𨸇𨸉𨸊𨸋𨸌𨸎𨸘𨸟𩏼𩏽𩏾𩏿𩐀𩓋𩖕𩖖𩖗𩙥𩙦𩙧𩙨𩙩𩙪𩙫𩙬𩙭𩙮𩙯𩙰𩟿𩠀𩠁𩠂𩠃𩠅𩠆𩠇𩠈𩠉𩠊𩠋𩠌𩠎𩠏𩠠𩡖𩧦𩧨𩧩𩧪𩧫𩧬𩧭𩧮𩧯𩧰𩧱𩧲𩧳𩧴𩧵𩧶𩧸𩧺𩧻𩧼𩧿𩨀𩨁𩨂𩨃𩨄𩨅𩨆𩨇𩨈𩨉𩨊𩨋𩨌𩨍𩨎𩨏𩨐𩩈𩬣𩬤𩭹𩯒𩰰𩲒𩴌𩽹𩽺𩽻𩽼𩽽𩽾𩽿𩾁𩾂𩾃𩾄𩾅𩾆𩾇𩾈𩾊𩾋𩾌𩾎𪉂𪉃𪉄𪉅𪉆𪉈𪉉𪉊𪉋𪉌𪉍𪉎𪉏𪉐𪉑𪉒𪉔𪉕𪎈𪎉𪎊𪎋𪎌𪑅𪔭𪚏𪚐𪜎𪞝𪟎𪟝𪠀𪠟𪠡𪠳𪠵𪠸𪠺𪠽𪡀𪡃𪡋𪡏𪡛𪡞𪡺𪢌𪢐𪢒𪢕𪢖𪢠𪢮𪢸𪣆𪣒𪣻𪤄𪤚𪥠𪥫𪥰𪥿𪧀𪧘𪨊𪨗𪨧𪨩𪨶𪨷𪨹𪩇𪩎𪩘𪩛𪩷𪩸𪪏𪪑𪪞𪪴𪪼𪫌𪫡𪫷𪫺𪬚𪬯𪭝𪭢𪭧𪭯𪭵𪭾𪮃𪮋𪮖𪮳𪮶𪯋𪰶𪱥𪱷𪲎𪲔𪲛𪲮𪳍𪳗𪴙𪵑𪵣𪵱𪶄𪶒𪶮𪷍𪷽𪸕𪸩𪹀𪹠𪹳𪹹𪺣𪺪𪺭𪺷𪺸𪺻𪺽𪻐𪻨𪻲𪻺𪼋𪼴𪽈𪽝𪽪𪽭𪽮𪽴𪽷𪾔𪾢𪾣𪾦𪾸𪿊𪿞𪿫𪿵𫀌𫀓𫀨𫀬𫀮𫁂𫁟𫁡𫁱𫁲𫁳𫁷𫁺𫂃𫂆𫂈𫂖𫂿𫃗𫄙𫄚𫄛𫄜𫄝𫄞𫄟𫄠𫄡𫄢𫄣𫄤𫄥𫄦𫄧𫄨𫄩𫄪𫄫𫄬𫄭𫄮𫄯𫄰𫄱𫄲𫄳𫄴𫄵𫄶𫄷𫄸𫄹𫅅𫅗𫅥𫅭𫅼𫆏𫆝𫆫𫇘𫇛𫇪𫇭𫇴𫇽𫈉𫈎𫈟𫈵𫉁𫉄𫊪𫊮𫊸𫊹𫊻𫋇𫋌𫋲𫋷𫋹𫋻𫌀𫌇𫌋𫌨𫌪𫌫𫌬𫌭𫌯𫍐𫍙𫍚𫍛𫍜𫍝𫍞𫍟𫍠𫍡𫍢𫍣𫍤𫍥𫍦𫍧𫍨𫍩𫍪𫍫𫍬𫍭𫍮𫍯𫍰𫍱𫍲𫍳𫍴𫍵𫍶𫍷𫍸𫍹𫍺𫍻𫍼𫍽𫍾𫍿𫎆𫎌𫎦𫎧𫎨𫎩𫎪𫎫𫎬𫎭𫎱𫎳𫎸𫎺𫏃𫏆𫏋𫏌𫏐𫏑𫏕𫏞𫏨𫐄𫐅𫐆𫐇𫐈𫐉𫐊𫐋𫐌𫐍𫐎𫐏𫐐𫐑𫐒𫐓𫐔𫐕𫐖𫐗𫐘𫐙𫐷𫑘𫑡𫑷𫓥𫓦𫓧𫓨𫓩𫓪𫓫𫓬𫓭𫓮𫓯𫓰𫓱𫓲𫓳𫓴𫓵𫓶𫓷𫓸𫓹𫓺𫓻𫓼𫓽𫓾𫓿𫔀𫔁𫔂𫔃𫔄𫔅𫔆𫔇𫔈𫔉𫔊𫔋𫔌𫔍𫔎𫔏𫔐𫔑𫔒𫔓𫔔𫔕𫔖𫔭𫔮𫔯𫔰𫔲𫔴𫔵𫔶𫔽𫕚𫕥𫕨𫖃𫖅𫖇𫖑𫖒𫖓𫖔𫖕𫖖𫖪𫖫𫖬𫖭𫖮𫖯𫖰𫖱𫖲𫖳𫖴𫖵𫖶𫖷𫖸𫖹𫖺𫗇𫗈𫗉𫗊𫗋𫗚𫗞𫗟𫗠𫗡𫗢𫗣𫗤𫗥𫗦𫗧𫗨𫗩𫗪𫗫𫗬𫗭𫗮𫗯𫗰𫗱𫗳𫗴𫗵𫘛𫘜𫘝𫘞𫘟𫘠𫘡𫘣𫘤𫘥𫘦𫘧𫘨𫘩𫘪𫘫𫘬𫘭𫘮𫘯𫘰𫘱𫘽𫙂𫚈𫚉𫚊𫚋𫚌𫚍𫚎𫚏𫚐𫚑𫚒𫚓𫚔𫚕𫚖𫚗𫚘𫚙𫚚𫚛𫚜𫚝𫚞𫚟𫚠𫚡𫚢𫚣𫚤𫚥𫚦𫚧𫚨𫚩𫚪𫚫𫚬𫚭𫛚𫛛𫛜𫛝𫛞𫛟𫛠𫛡𫛢𫛣𫛤𫛥𫛦𫛧𫛨𫛩𫛪𫛫𫛬𫛭𫛮𫛯𫛰𫛱𫛲𫛳𫛴𫛵𫛶𫛷𫛸𫛹𫛺𫛻𫛼𫛽𫛾𫜀𫜁𫜂𫜃𫜄𫜅𫜊𫜑𫜒𫜓𫜔𫜕𫜙𫜟𫜨𫜩𫜪𫜫𫜬𫜭𫜮𫜯𫜰𫜲𫜳𫝈𫝋𫝦𫝧𫝨𫝩𫝪𫝫𫝬𫝭𫝮𫝵𫞅𫞗𫞚𫞛𫞝𫞠𫞡𫞢𫞣𫞥𫞦𫞧𫞨𫞩𫞷𫟃𫟄𫟅𫟆𫟇𫟑𫟕𫟞𫟟𫟠𫟡𫟢𫟤𫟥𫟦𫟫𫟬𫟲𫟳𫟴𫟵𫟶𫟷𫟸𫟹𫟺𫟻𫟼𫟽𫟾𫟿𫠀𫠁𫠂𫠅𫠆𫠇𫠈𫠊𫠋𫠌𫠏𫠐𫠑𫠒𫠖𫠜𫢸𫧃𫧮𫫇𫬐𫭟𫭢𫭼𫮃𫰛𫵷𫶇𫷷𫸩𬀩𬀪𬂩𬃊𬇕𬇙𬇹𬉼𬊈𬊤𬍛𬍡𬍤𬒈𬒗𬕂𬘓𬘘𬘡𬘩𬘫𬘬𬘭𬘯𬙂𬙊𬙋𬜬𬜯𬞟𬟁𬟽𬣙𬣞𬣡𬣳𬤇𬤊𬤝𬨂𬨎𬩽𬪩𬬩𬬭𬬮𬬱𬬸𬬹𬬻𬬿𬭁𬭊𬭎𬭚𬭛𬭤𬭩𬭬𬭭𬭯𬭳𬭶𬭸𬭼𬮱𬮿𬯀𬯎𬱖𬱟𬳵𬳶𬳽𬳿𬴂𬴃𬴊𬶋𬶍𬶏𬶐𬶟𬶠𬶨𬶭𬶮𬷕𬸘𬸚𬸣𬸦𬸪𬸯𬹼𬺈𬺓𰬸𰰨𰶎𰾄𰾭𱊜")
def rgb(h): return RGBColor.from_string(h)

# 圖示：用 PowerPoint 原生圖形畫，任何電腦都畫得出來，也能在 PowerPoint 裡點選修改。
# 值是（圖形, 圖形上的字）；"bars" 是三根長條。deck.json 每頁可寫 "icon": "名稱"；寫 "none" 就不放。
ICONS = {
    "speed": (MSO_SHAPE.LIGHTNING_BOLT, None),   "growth": (MSO_SHAPE.UP_ARROW, None),
    "decline": (MSO_SHAPE.DOWN_ARROW, None),     "compare": (MSO_SHAPE.LEFT_RIGHT_ARROW, None),
    "data": (MSO_SHAPE.CAN, None),               "method": (MSO_SHAPE.GEAR_6, None),
    "paper": (MSO_SHAPE.FOLDED_CORNER, None),    "key": (MSO_SHAPE.STAR_5_POINT, None),
    "target": (MSO_SHAPE.DONUT, None),           "chart": ("bars", None),
    "warning": (MSO_SHAPE.ISOSCELES_TRIANGLE, "!"), "question": (MSO_SHAPE.OVAL, "?"),
    "money": (MSO_SHAPE.OVAL, "$"),              "percent": (MSO_SHAPE.OVAL, "%"),
    "check": (MSO_SHAPE.OVAL, "✓"),
}
# 沒寫 icon 時依標題的字自動挑（由上往下，第一個對到的）；都對不到就看版面
AUTO_ICON = [("key", r"^(結論|結語|總結)|重點|啟示"), ("question", r"[？?]|為何|為什麼|如何"), ("speed", r"加速|速度|毫秒|微秒|即時|更快"),
             ("decline", r"下降|降低|減少|縮小|低估|下跌"), ("warning", r"風險|限制|誤差|偏誤|挑戰|瓶頸|洩漏|失效"),
             ("growth", r"提升|成長|增加|上升|擴大|優於|高於|超過"), ("data", r"資料|樣本|數據|筆數|資料集"),
             ("method", r"模型|方法|架構|演算法|神經網路|校準|訓練"), ("paper", r"文獻|論文|期刊|綜述|研究"),
             ("money", r"價格|定價|報酬|成本|獲利|交易|投資")]
LAYOUT_ICON = {"bignumber": "chart", "compare": "compare", "table": "data", "quote": "paper", "assertion": "key", "section": "key"}

def icon_of(sl, kind):
    name = sl.get("icon")
    if name == "none": return None
    if name in ICONS: return name
    if kind == "section":   # 章節頁（開場、結語）不依關鍵字挑，免得「開場：…風險管理」配上警告標誌
        return "question" if re.search(r"[？?]", sl.get("title", "")) else "key"
    for k, pat in AUTO_ICON:
        if re.search(pat, sl.get("title", "")): return k
    return LAYOUT_ICON.get(kind)

def draw_icon(slide, th, name, l, t, size):
    shape, glyph = ICONS[name]
    if shape == "bars":   # 長條圖：三根由低到高
        w = size / 4.2
        for i, h in enumerate((0.42, 0.66, 0.95)):
            b = box(slide, l + i * (w * 1.4), t + size * (1 - h), w, size * h, th["accent"])
        return
    if shape == MSO_SHAPE.LEFT_RIGHT_ARROW:   # 方框裡的左右箭頭會擠成菱形，改成橫的
        t, size_h = t + size * 0.2, size * 0.6
    else:
        size_h = size
    s = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(size), Inches(size_h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(th["accent"]); s.line.fill.background(); s.shadow.inherit = False
    if shape == MSO_SHAPE.FOLDED_CORNER:      # 文件：加三條「文字線」，不然看起來只是方塊
        for i in range(3):
            box(slide, l + size * 0.18, t + size * (0.25 + i * 0.2), size * (0.64 if i < 2 else 0.4), size * 0.07, th["bg"])
    if glyph:
        tf = s.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        tri = shape == MSO_SHAPE.ISOSCELES_TRIANGLE   # 三角形上窄下寬：字放在底部才不會被斜邊切掉
        if tri: tf.vertical_anchor = MSO_ANCHOR.BOTTOM; tf.margin_bottom = Inches(size * 0.04)
        r = p.add_run(); r.text = glyph; style(r, int(size * (30 if tri else 40)), True, th["bg"])

def title_lines(title, width, t_size):
    """估算標題會折成幾行（中文字算兩格）。"""
    units = sum(2 if ord(ch) > 0x2E7F else 1 for ch in title)
    per_line = 2 * width * 72 / t_size * 0.92
    return max(1, -(-int(units) // int(per_line)))
W, H = 13.333, 7.5
THEMES = {  # 三套風格：配色、字級、裝飾都不同，內容不變
    "academic":  dict(name="學術簡潔", bg="FFFFFF", ink="1C2024", muted="6B7280", title="0B3D6B", accent="B3261E",
                      soft="EDF1F6", row="F6F7F9", t_size=30, b_size=22, cover=40),
    "editorial": dict(name="編輯雜誌", bg="F7F3EC", ink="1A1A1A", muted="6F675C", title="1A1A1A", accent="C2410C",
                      soft="EADFCF", row="F1EADF", t_size=34, b_size=22, cover=48),
    "stage":     dict(name="深色講台", bg="0F172A", ink="F1F5F9", muted="94A3B8", title="FFFFFF", accent="F59E0B",
                      soft="1E293B", row="16213A", t_size=36, b_size=24, cover=50),
}

def style(run, size, bold=False, color="000000"):
    f = run.font; f.size = Pt(size); f.bold = bold; f.color.rgb = rgb(color); f.name = LATIN
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {}); rPr.append(el)
        el.set("typeface", EA)

def plain(x):
    """模型常把 Markdown 的 **粗體** 記號寫進來，PowerPoint 不認得，會原樣印出星號，所以先拿掉。"""
    return re.sub(r"\*\*|__|`", "", str(x))

def text(slide, l, t, w, h, lines, size, color, bold=False, bullet=None, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=10):
    tf = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h)).text_frame
    tf.word_wrap = True; tf.vertical_anchor = anchor
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        if bullet:  # 真正的項目符號＋懸掛縮排，折行時第二行會對齊文字而不是對齊符號
            pPr = p._p.get_or_add_pPr(); pPr.set("marL", str(Inches(0.35))); pPr.set("indent", str(-Inches(0.35)))
            clr = pPr.makeelement(qn("a:buClr"), {}); c = clr.makeelement(qn("a:srgbClr"), {"val": bullet}); clr.append(c); pPr.append(clr)
            pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
            pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "■" if size < 30 else "•"}))
        r = p.add_run(); r.text = plain(line); style(r, size, bold, color)

def box(slide, l, t, w, h, color):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(color); s.line.fill.background(); s.shadow.inherit = False
    return s

def new_slide(prs, th, bg=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(bg or th["bg"])
    return s

def header(s, th, key, sl, icon=None):
    """標題區：小標（可省略）＋一句結論當標題＋風格裝飾＋右上角圖示。回傳內容區的起始高度。"""
    left = 1.05 if key == "editorial" else 0.8
    width = 11.7 - (1.15 if icon else 0)
    if key == "editorial":
        box(s, 0.45, 0.55, 0.14, 1.35, th["accent"])
    if icon:
        draw_icon(s, th, icon, W - 0.8 - 0.85, 0.72, 0.85)
    if sl.get("kicker"):
        text(s, left, 0.42, width, 0.35, sl["kicker"], 13, th["accent"], True)
    text(s, left, 0.72, width, 1.25, sl["title"], th["t_size"], th["title"], True, anchor=MSO_ANCHOR.TOP)
    # 標題折成多行時，把下面的內容依行數往下推，才不會壓到標題（原本只處理到兩行）
    extra = 0.6 * (min(title_lines(sl["title"], width, th["t_size"]), 4) - 1)
    if key != "editorial":
        box(s, left, 1.95 + extra, 1.5, 0.045, th["accent"])
    return 2.3 + extra

def footer(s, th, sl, n, total):
    if sl.get("note"):
        text(s, 0.8, 6.92, 10.6, 0.4, sl["note"], 11, th["muted"])
    text(s, 11.5, 6.92, 1.1, 0.4, f"{n} / {total}", 11, th["muted"], align=PP_ALIGN.RIGHT)

def bullets(s, th, l, t, w, items, size=None):
    text(s, l, t, w, 4.4, items, size or th["b_size"], th["ink"], bullet=th["accent"], space=14)

def add_table(s, th, rows, top):
    n, m = len(rows), max(len(r) for r in rows)
    size = 20 if n <= 5 else (16 if n <= 7 else (13 if n <= 10 else 11))   # 列數少就放大，不要讓小表格縮在角落
    tbl = s.shapes.add_table(n, m, Inches(0.8), Inches(top), Inches(11.7), Inches(min(4.4, (0.7 if n <= 5 else 0.42) * n))).table
    width = lambda x: sum(2 if ord(ch) > 0x2E7F else 1 for ch in str(x))   # 中文字算兩格
    longest = [max(width(r[j]) if j < len(r) else 0 for r in rows) for j in range(m)]
    share = [max(x, 0.1 * sum(longest)) for x in longest]
    for j in range(m):
        tbl.columns[j].width = Inches(11.7 * share[j] / sum(share))
    for i, row in enumerate(rows):
        for j in range(m):
            c = tbl.cell(i, j); c.text = ""
            r = c.text_frame.paragraphs[0].add_run(); r.text = plain(row[j]) if j < len(row) else ""
            style(r, size, i == 0, th["title"] if i == 0 else th["ink"])
            c.fill.solid(); c.fill.fore_color.rgb = rgb(th["soft"] if i == 0 else (th["bg"] if i % 2 else th["row"]))

def layout_of(sl):
    if sl.get("layout"): return sl["layout"]
    return "table" if sl.get("table") else "assertion"

def build(deck, out):
    key = deck.get("theme", "academic"); th = THEMES[key]
    prs = Presentation(); prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    total = len(deck["slides"]) + 1
    # 封面
    s = new_slide(prs, th)
    if key == "editorial":
        box(s, 0, 0, 0.5, H, th["accent"])
    elif key == "stage":
        box(s, 0.9, 2.2, 0.9, 0.09, th["accent"])
    text(s, 1.0 if key == "editorial" else 0.9, 2.45, 11.3, 2.0, deck["title"], th["cover"], th["title"], True)
    if deck.get("subtitle"):
        text(s, 1.0 if key == "editorial" else 0.9, 4.55, 11.3, 0.9, deck["subtitle"], 20, th["muted"])
    if key == "academic":
        box(s, 0.9, 4.35, 1.6, 0.05, th["accent"])
    # 內容頁
    for n, sl in enumerate(deck["slides"], 2):
        kind = layout_of(sl)
        if kind == "section":   # 章節頁：整頁只有一句話，讓聽眾換氣
            s = new_slide(prs, th, th["soft"]); icon = icon_of(sl, kind)
            box(s, 0.9, 3.0, 0.12, 1.5, th["accent"])
            if icon: draw_icon(s, th, icon, W - 0.9 - 1.4, 3.05, 1.4)
            if sl.get("kicker"): text(s, 1.3, 2.55, 10.5, 0.4, sl["kicker"], 14, th["accent"], True)
            text(s, 1.3, 3.0, 10.8 - (1.6 if icon else 0), 1.6, sl["title"], th["cover"] - 6, th["title"], True, anchor=MSO_ANCHOR.MIDDLE)
            if sl.get("subtitle"): text(s, 1.3, 4.7, 10.8, 0.8, sl["subtitle"], 18, th["muted"])
            footer(s, th, sl, n, total); continue
        s = new_slide(prs, th); top = header(s, th, key, sl, icon_of(sl, kind))
        if kind == "bignumber":  # 大數字：一個數字撐起一頁
            num = plain(sl["number"]); units = sum(2 if ord(ch) > 0x2E7F else 1 for ch in num)
            big = 110 if units <= 4 else (80 if units <= 7 else (60 if units <= 10 else 44))   # 字數多就縮小，保持在一行內
            text(s, 0.8, top + 0.1, 5.2, 2.0, num, big, th["accent"], True, anchor=MSO_ANCHOR.MIDDLE)
            if sl.get("caption"): text(s, 0.8, top + 2.3, 5.2, 1.2, sl["caption"], 18, th["muted"])
            if sl.get("bullets"):
                box(s, 6.25, top + 0.2, 0.03, 3.6, th["soft"]); bullets(s, th, 6.6, top + 0.2, 5.9, sl["bullets"], th["b_size"] - 2)
        elif kind == "table":
            add_table(s, th, sl["table"], top - 0.1)
            if sl.get("bullets"): bullets(s, th, 0.9, 6.0, 11.5, sl["bullets"], 14)
        elif kind == "compare":  # 兩欄對照
            for col, x in ((sl["left"], 0.8), (sl["right"], 6.95)):
                box(s, x, top, 5.6, 0.62, th["soft"])
                text(s, x + 0.2, top + 0.08, 5.2, 0.5, col["heading"], 20, th["title"], True, anchor=MSO_ANCHOR.MIDDLE)
                bullets(s, th, x + 0.1, top + 0.85, 5.4, col.get("bullets") or [], th["b_size"] - 3)
        elif kind == "quote":    # 引言：原文照抄，下面註明出處
            text(s, 0.8, top - 0.2, 1.2, 1.2, "“", 96, th["accent"], True)
            text(s, 1.7, top + 0.25, 10.6, 2.8, sl["quote"], 28, th["ink"], anchor=MSO_ANCHOR.TOP)
            if sl.get("by"): text(s, 1.7, top + 3.2, 10.6, 0.5, "— " + sl["by"], 16, th["muted"])
        else:                    # assertion：一句結論＋支持它的證據
            bullets(s, th, 0.9 if key != "editorial" else 1.1, top, 11.4, sl.get("bullets") or [])
        footer(s, th, sl, n, total)
    prs.save(out)
    return prs, th

# ---------- 檢查 ----------
def slide_texts(sl):
    out = [sl.get("title", ""), sl.get("note", ""), sl.get("kicker", ""), sl.get("subtitle", ""),
           sl.get("number", ""), sl.get("caption", ""), sl.get("quote", ""), sl.get("by", "")] + list(sl.get("bullets") or [])
    for row in sl.get("table") or []:
        out += [str(c) for c in row]
    for side in ("left", "right"):
        if sl.get(side):
            out += [sl[side].get("heading", "")] + list(sl[side].get("bullets") or [])
    return [str(x) for x in out if x]

DASH = re.compile(r"(?<=\d)\s*(?:[-‐‑‒–—~～]|至|到)\s*(?=\d)")   # 數字間的各種連字號（含 U+2011）與「至」「到」統一成 -

def facts(t):
    """抽出需要有來源的東西：DOI、ISSN、數字範圍（2‑3%、9,000–16,000）、帶 % 的數字（個位數也算）、兩位數以上的數字"""
    doi = r"10\.\d{4,9}/[^\s|)\]，。；]+"; issn = r"\d{4}-\d{3}[\dXx]"
    found = set(re.findall(doi, t)); rest = re.sub(doi, " ", t)
    found |= set(re.findall(issn, rest)); rest = re.sub(issn, " ", rest)
    rng = r"\d[\d,]*\.?\d*\s*(?:[-‐‑‒–—~～]|至|到)\s*\d[\d,]*\.?\d*\s*[%％]?"   # 範圍要整段對上，不能拆成兩個數字各自過關
    found |= {DASH.sub("-", x).replace("％", "%").replace(" ", "") for x in re.findall(rng, rest)}; rest = re.sub(rng, " ", rest)
    found |= {n.replace(" ", "") for n in re.findall(r"\d[\d,]*\.?\d*\s*[%％]", rest)}; rest = re.sub(r"\d[\d,]*\.?\d*\s*[%％]", " ", rest)
    found |= {n for n in re.findall(r"\d[\d,]*\.?\d*%?", rest) if len(re.sub(r"\D", "", n)) >= 2}
    return found

UNITS = [("ms", "毫秒"), ("µs", "μs", "us", "微秒"), ("秒",), ("%", "％", "percent"), ("倍",), ("篇",), ("年",),
         ("筆",), ("組",), ("bps", "基點"), ("個",), ("天", "日"), ("億",), ("萬",)]

def bignumber_ok(sl, blob):
    """大數字的每個數字都要在來源找得到；兩位數以下的還要和單位一起出現（單位取自 number 與 caption）"""
    num = re.sub(r"(?<=\d),(?=\d)", "", plain(sl["number"]))
    blob = re.sub(r"(?<=\d)[,，](?=\d)", "", blob)
    ctx = num + " " + plain(sl.get("caption", ""))
    units = [g for g in UNITS if any(u.lower() in ctx.lower() for u in g)]
    for n in re.findall(r"\d+(?:\.\d+)?(?![\d.]*[A-Za-z])", num):   # 「1F Bergomi」「3D」的數字是名稱的一部分，不算
        whole = r"(?<![\d.])" + re.escape(n) + r"(?!\d|\.\d)"
        if len(re.sub(r"\D", "", n)) >= 2:
            if not re.search(whole, blob): return False
        elif units:
            pat = whole + r"\s*(?:" + "|".join(re.escape(u) for g in units for u in g) + ")"
            if not re.search(pat, blob, re.I): return False
        elif not re.search(whole, blob): return False
    return True

def check_sources(deck, base):
    srcs = deck.get("sources") or []
    if not srcs:
        return None, "deck.json 沒有填 sources，無法核對來源"
    blob = ""
    for p in srcs:
        if not (base / p).exists():
            return None, f"來源檔不存在：{p}"
        blob += (base / p).read_text(encoding="utf-8", errors="replace") + "\n"
    missing = []
    for i, sl in enumerate([{"title": deck.get("title", ""), "subtitle": deck.get("subtitle", "")}] + deck["slides"]):
        found = set()
        for t in slide_texts(sl):
            found |= not_in_source(t, blob)
        # 大數字頁的數字是整頁最顯眼的主張：個位數也要核對，而且要連單位一起對上
        #（來源寫「約 6.7 ms」卻寫成「1」毫秒；只比對「1」的話，來源裡的「第 1 頁」「1F」都會讓它過關）
        if layout_of(sl) == "bignumber" and sl.get("number"):
            if not bignumber_ok(sl, blob): found.add(plain(sl["number"]))
        missing += [f"{'封面' if i == 0 else f'第 {i} 頁'}的 {x}" for x in sorted(found)]
    return missing, None

def not_in_source(text, blob):
    """回傳 text 裡在來源找不到的數字、DOI、ISSN。數字要整個對上：30 不能靠 30.9 或 300,000 過關；有 % 的要連 % 一起對上。"""
    blob = DASH.sub("-", re.sub(r"(?<=\d)[,，](?=\d)", "", blob)).replace("％", "%").lower(); out = set()
    for x in facts(text):
        k = re.sub(r"(?<=\d),(?=\d)", "", x).rstrip(".").lower()
        if re.fullmatch(r"[\d.]+-[\d.]+%?", k):   # 範圍：兩端數字與 % 都要在同一處出現
            a, b = k.rstrip("%").split("-")
            pat = r"(?<![\d.])" + re.escape(a) + r"-" + re.escape(b) + r"(?!\d|\.\d)" + (r"\s*%" if k.endswith("%") else "")
            whole = lambda n: re.search(r"(?<![\d.])" + re.escape(n) + r"(?!\d|\.\d)", blob)
            # 來源把兩端分開寫（「300,000 µs（1F）至 500,000 µs（rB）」）時，兩端都是兩位數以上、各自找得到也算；
            # 個位數或帶 % 的範圍（「2‑3%」）太容易湊到頁碼之類的數字，一定要整段對上
            loose = not k.endswith("%") and min(len(re.sub(r"\D", "", a)), len(re.sub(r"\D", "", b))) >= 2 and whole(a) and whole(b)
            if not re.search(pat, blob) and not loose: out.add(x)
        elif re.fullmatch(r"[\d.]+%?", k):
            pat = r"(?<![\d.])" + re.escape(k.rstrip("%")) + r"(?!\d|\.\d)" + (r"\s*[%％]" if k.endswith("%") else "")
            if not re.search(pat, blob): out.add(x)
        elif k not in blob:
            out.add(x)
    return out

def check_outline(deck, base):
    f = base / "outline.md"
    if not f.exists():
        return ["找不到 outline.md：要先寫大綱、給使用者確認，才能做投影片"]
    norm = lambda x: re.sub(r"[\s*`｜|]", "", x)
    outline = norm(f.read_text(encoding="utf-8"))
    return [f"第 {i} 頁標題「{sl['title']}」不在 outline.md 裡" for i, sl in enumerate(deck["slides"], 1) if norm(sl["title"]) not in outline]

def check_fonts(prs):
    total = bad = 0
    for s in prs.slides:
        for r in s._element.iter(qn("a:r")):
            total += 1
            rPr = r.find(qn("a:rPr")); ea = rPr.find(qn("a:ea")) if rPr is not None else None
            bad += ea is None or not ea.get("typeface")
    return total, bad

def check_language(deck):
    """找出簡體字：回傳 [(第幾頁, 字)]"""
    hits = []
    for i, sl in enumerate([{"title": deck.get("title", ""), "subtitle": deck.get("subtitle", "")}] + deck["slides"]):
        found = sorted({ch for t in slide_texts(sl) for ch in t if ch in SIMPLIFIED})
        if found: hits.append(("封面" if i == 0 else f"第 {i} 頁", "".join(found)))
    return hits

# 「具體」＝有數字、英文名稱（模型名、方法名、作者）或引號裡的原文
CONCRETE = re.compile(r"\d|「[^」]+」")
def concrete(x, title):
    """要點有沒有具體事實：數字、「」引文，或標題裡沒出現過的英文名稱（「ANN 表現較佳」裡的 ANN 已在標題，不算）"""
    return bool(CONCRETE.search(x)) or any(w.lower() not in title.lower() for w in re.findall(r"[A-Za-z][A-Za-z\-]{2,}", x))

def check_size(deck):
    warn, run, prev = [], 0, None
    for i, sl in enumerate(deck["slides"], 1):
        kind = layout_of(sl); run = run + 1 if kind == prev else 1; prev = kind
        if run == MAX_SAME_LAYOUT + 1 and kind != "section":
            warn.append(f"第 {i} 頁起連續 {run} 頁都是 {kind} 版面，換一種版面或插一頁章節頁")
        if kind == "section" and run == MAX_SECTIONS + 1:
            warn.append(f"第 {i} 頁是連續第 {run} 頁章節頁（只有標題），太空了：把其中一頁改成 assertion，寫出這段要聽眾記住的事實")
        if kind == "section" and any(sl.get(k) for k in ("bullets", "table", "left", "right", "quote", "number")):
            warn.append(f"第 {i} 頁是章節頁，章節頁只會顯示標題（與 kicker、subtitle），寫在上面的要點不會出現：改成 assertion 版面，或把要點移到別頁")
        if GENERIC.match(sl["title"].strip()):
            warn.append(f"第 {i} 頁標題「{sl['title']}」太籠統，請改成一句結論或一個問題")
        if sl.get("icon") and sl["icon"] != "none" and sl["icon"] not in ICONS:
            warn.append(f"第 {i} 頁的 icon「{sl['icon']}」不存在，可用的是 {' / '.join(ICONS)}（或 none）")
        width = 10.55 if kind != "section" else 9.2
        if title_lines(sl["title"], width, 34) > 2:
            warn.append(f"第 {i} 頁標題太長（會折成 3 行以上、壓到內容）：{sl['title'][:20]}…，請縮短到兩行內（大綱要一起改）")
        groups = [sl.get("bullets") or []] + [sl[k].get("bullets") or [] for k in ("left", "right") if sl.get(k)]
        flat = [x for b in groups for x in b]
        if kind == "assertion" and len(sl.get("bullets") or []) < 2:
            warn.append(f"第 {i} 頁只有 {len(sl.get('bullets') or [])} 條要點：assertion 至少要 2 條證據支持標題")
        if kind in ("assertion", "compare") and flat and not any(concrete(x, sl["title"]) for x in flat):
            warn.append(f"第 {i} 頁的要點都是形容詞，沒有具體事實：至少一條要放來源裡的數字、年份、樣本數或名稱（原樣抄）")
        for b in groups:
            if len(b) > MAX_BULLETS:
                warn.append(f"第 {i} 頁有 {len(b)} 條要點（上限 {MAX_BULLETS}）")
            warn += [f"第 {i} 頁有一條 {len(x)} 字（上限 {MAX_CHARS}）：{x[:18]}…" for x in b if len(x) > MAX_CHARS]
    return warn

def validate(deck):
    for k in ("title", "slides"):
        if k not in deck: return f"deck.json 缺少 '{k}'"
    if deck.get("theme", "academic") not in THEMES:
        return f"theme 必須是 {' / '.join(THEMES)} 其中之一"
    need = {"bignumber": ["number"], "table": ["table"], "compare": ["left", "right"], "quote": ["quote"]}
    for i, sl in enumerate(deck["slides"], 1):
        if "title" not in sl: return f"第 {i} 頁缺少 'title'"
        kind = layout_of(sl)
        if kind not in LAYOUTS: return f"第 {i} 頁的 layout '{kind}' 不存在，可用的是 {' / '.join(LAYOUTS)}"
        for k in need.get(kind, []):
            if not sl.get(k): return f"第 {i} 頁是 {kind} 版面，缺少 '{k}'"
        if kind == "assertion" and not sl.get("bullets"): return f"第 {i} 頁是 assertion 版面，缺少 'bullets'"
    return None

def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    src = pathlib.Path(args[0] if args else "deck.json"); out = args[1] if len(args) > 1 else "slides.pptx"
    try:
        deck = json.loads(src.read_text(encoding="utf-8"))
    except Exception as e:
        sys.exit(f"BUILD FAILED: 讀不了 {src}：{e}")
    err = validate(deck)
    if err: sys.exit(f"BUILD FAILED: {err}")
    prs, th = build(deck, out); base = src.resolve().parent; fail = False
    print(f"BUILD OK: {out}（{len(prs.slides)} 頁，含封面；風格 {th['name']}；中文字型 {EA}）")
    total, bad = check_fonts(prs)
    print(f"[字型] {total} 段文字，{total - bad} 段有東亞字型" + ("" if not bad else f" ← {bad} 段沒有，中文會跑版"))
    miss = check_outline(deck, base); fail |= bool(miss)
    print("[大綱] " + ("通過：每一頁的標題都出自 outline.md" if not miss else "；".join(miss)))
    warn = check_size(deck)
    print("[份量] " + ("通過" if not warn else "；".join(warn)))
    zh = check_language(deck); fail |= bool(zh)
    print("[語言] " + ("通過：沒有簡體字" if not zh else "有簡體字，請改成繁體：" + "；".join(f"{w}「{c}」" for w, c in zh)))
    missing, err = check_sources(deck, base)
    if err:
        print("[來源] " + err); fail = True
    elif missing:
        print(f"[來源] 有 {len(missing)} 項在來源檔裡找不到：{'、'.join(missing)}")
        print("       → 回來源檔確認；是寫錯就改 deck.json，是自己加的就刪掉，然後重跑。"); fail = True
    else:
        print("[來源] 通過：投影片裡的數字、DOI、ISSN 都能在來源檔找到")
    if "--pdf" in sys.argv:
        exe = shutil.which("soffice") or shutil.which("libreoffice")
        if exe:
            subprocess.run([exe, "--headless", "--convert-to", "pdf", out], capture_output=True, timeout=180)
            print(f"[預覽] 已另存 {pathlib.Path(out).with_suffix('.pdf')}")
        else:
            print("[預覽] 這台電腦沒有 LibreOffice，請直接用 PowerPoint 開啟 " + out)
    sys.exit(2 if fail else 0)

if __name__ == "__main__":
    main()
